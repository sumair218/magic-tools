import csv
import os
import shutil
import uuid

from django.conf import settings
from django.contrib import messages
from django.http import FileResponse
from django.shortcuts import render

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

try:
    import pytesseract
except ImportError:
    pytesseract = None

from PIL import Image, ImageEnhance, ImageFilter, ImageOps

from .forms import OCRImageForm, OCRPDFForm


def _resolve_tesseract_cmd():
    if not pytesseract:
        return None

    explicit_cmd = getattr(settings, 'TESSERACT_CMD', None)
    if explicit_cmd:
        return explicit_cmd

    path = shutil.which('tesseract')
    if path:
        return path

    windows_candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for candidate in windows_candidates:
        if os.path.exists(candidate):
            return candidate

    return None


def _save_temp_file(uploaded_file):
    temp_dir = settings.MEDIA_ROOT / 'temp'
    temp_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{uuid.uuid4().hex}_{uploaded_file.name}"
    file_path = temp_dir / filename
    with open(file_path, 'wb+') as destination:
        for chunk in uploaded_file.chunks():
            destination.write(chunk)
    return file_path


def _serve_text_file(path, filename):
    response = FileResponse(
        open(path, 'rb'),
        as_attachment=True,
        filename=filename
    )
    return response


def _save_text_as_pdf(text, path):
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    y = height - 72
    line_height = 14

    for line in text.splitlines():
        if y < 72:
            c.showPage()
            y = height - 72
        c.drawString(72, y, line)
        y -= line_height

    c.save()


def _save_text_as_docx(text, path):
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(path)


def _save_text_as_csv(text, path):
    with open(path, 'w', encoding='utf-8', newline='') as csvfile:
        writer = csv.writer(csvfile)
        for line in text.splitlines():
            writer.writerow([line])


def _save_text_as_xlsx(text, path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'OCR Text'
    for row_index, line in enumerate(text.splitlines(), start=1):
        worksheet.cell(row=row_index, column=1, value=line)
    workbook.save(path)


def _enhance_low_quality_image(image):
    image = image.convert('L')
    image = image.filter(ImageFilter.MedianFilter(size=3))
    image = ImageOps.equalize(image)
    image = ImageOps.autocontrast(image)
    image = image.resize((image.width * 2, image.height * 2), Image.LANCZOS)
    image = image.filter(ImageFilter.UnsharpMask(radius=2, percent=200, threshold=3))
    image = image.filter(ImageFilter.DETAIL)
    enhancer = ImageEnhance.Sharpness(image)
    image = enhancer.enhance(2.0)
    enhancer = ImageEnhance.Contrast(image)
    image = enhancer.enhance(1.7)
    enhancer = ImageEnhance.Brightness(image)
    image = enhancer.enhance(1.2)
    image = image.point(lambda x: 0 if x < 140 else 255)
    return image


def home(request):
    return render(request, 'ocr_tools/home.html')


def image_to_text(request):
    form = OCRImageForm(request.POST or None, request.FILES or None)
    tesseract_cmd = _resolve_tesseract_cmd()
    ocr_available = pytesseract is not None and tesseract_cmd is not None

    if request.method == 'POST' and form.is_valid():
        if not ocr_available:
            messages.error(request, 'OCR is not available because Tesseract is not installed or the executable could not be found.')
        else:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
            image_path = _save_temp_file(form.cleaned_data['image_file'])
            try:
                with Image.open(image_path) as image:
                    if form.cleaned_data.get('enhance_low_quality'):
                        image = _enhance_low_quality_image(image)
                    text = pytesseract.image_to_string(image, config='--oem 3 --psm 6')

                    if not text.strip() and form.cleaned_data.get('enhance_low_quality'):
                        inverted = ImageOps.invert(image.convert('L'))
                        text = pytesseract.image_to_string(inverted, config='--oem 3 --psm 6')

                output_format = form.cleaned_data['output_format']
                filename_base = os.path.splitext(form.cleaned_data['image_file'].name)[0]
                temp_dir = settings.MEDIA_ROOT / 'temp'
                temp_dir.mkdir(parents=True, exist_ok=True)

                if output_format == 'txt':
                    output_path = temp_dir / f"{uuid.uuid4().hex}_{filename_base}.txt"
                    with open(output_path, 'w', encoding='utf-8') as output_file:
                        output_file.write(text)
                elif output_format == 'pdf':
                    output_path = temp_dir / f"{uuid.uuid4().hex}_{filename_base}.pdf"
                    _save_text_as_pdf(text, output_path)
                elif output_format == 'word':
                    output_path = temp_dir / f"{uuid.uuid4().hex}_{filename_base}.docx"
                    _save_text_as_docx(text, output_path)
                elif output_format == 'csv':
                    output_path = temp_dir / f"{uuid.uuid4().hex}_{filename_base}.csv"
                    _save_text_as_csv(text, output_path)
                elif output_format == 'xlsx':
                    output_path = temp_dir / f"{uuid.uuid4().hex}_{filename_base}.xlsx"
                    _save_text_as_xlsx(text, output_path)
                else:
                    output_path = temp_dir / f"{uuid.uuid4().hex}_{filename_base}.txt"
                    with open(output_path, 'w', encoding='utf-8') as output_file:
                        output_file.write(text)

                output_name = os.path.basename(output_path)
                download_url = settings.MEDIA_URL.rstrip('/') + '/temp/' + output_name
                return render(request, 'ocr_tools/tool_form.html', {
                    'tool_name': 'Image to Text',
                    'tool_description': 'Upload an image and extract text, then download as TXT, PDF, Word, CSV, or XLSX.',
                    'form': form,
                    'download_url': download_url,
                    'download_name': output_name,
                    'extracted_text': text,
                })
            except Exception:
                messages.error(request, 'Unable to extract text from the uploaded image.')

    if not ocr_available:
        messages.error(request, 'OCR is not available because Tesseract is not installed or the executable could not be found.')

    return render(request, 'ocr_tools/tool_form.html', {
        'tool_name': 'Image to Text',
        'tool_description': 'Upload an image and extract text, then download as TXT, PDF, Word, CSV, or XLSX.',
        'form': form,
    })


def pdf_to_text(request):
    form = OCRPDFForm(request.POST or None, request.FILES or None)

    if request.method == 'POST' and form.is_valid():
        pdf_path = _save_temp_file(form.cleaned_data['pdf_file'])
        try:
            reader = PdfReader(pdf_path)
            text_content = []

            for page in reader.pages:
                page_text = page.extract_text() or ''
                text_content.append(page_text)

            text = '\n\n'.join(text_content)
            output_path = settings.MEDIA_ROOT / 'temp' / f"{uuid.uuid4().hex}.txt"
            with open(output_path, 'w', encoding='utf-8') as output_file:
                output_file.write(text)

            return _serve_text_file(text, output_path)
        except Exception:
            messages.error(request, 'Unable to extract text from the uploaded PDF.')

    return render(request, 'ocr_tools/tool_form.html', {
        'tool_name': 'PDF to Text',
        'tool_description': 'Upload a PDF and extract selectable text from each page.',
        'form': form,
    })
